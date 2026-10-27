"""Importar a planilha da escola: alunos, turmas e responsáveis de uma vez.

Ideia do integrador Control iD (áudio de 27/10): com a planilha importada a secretaria só
precisa tirar a foto. Cada escola exporta de um jeito (ActiveSoft, planilha própria, Google
Sheets) — "Nome da mãe", "Celular 1", "Série/Turma"... Por isso o sistema adivinha o
mapeamento pelo cabeçalho e a secretaria confere numa tela antes de gravar.

Regras que evitam retrabalho:
- irmãos com a mesma mãe → uma pessoa só (casada pelo celular), vinculada aos dois;
- reimportar a mesma planilha não duplica (aluno por matrícula, ou nome + turma;
  responsável pelo celular, depois e-mail);
- tudo numa transação (db.lote): ou importa tudo, ou nada.
"""
import csv
import io
import re
import unicodedata
from collections import Counter
from datetime import date, datetime

from .. import db
from ..auditoria import auditar
from ..dispositivos import service as dispositivos
from ..util import mesmo_telefone, normalizar_telefone
from . import service

MAX_BYTES = 5 * 1024 * 1024
MAX_LINHAS = 5000
SLOTS = service.MAX_RESPONSAVEIS

CAMPOS_ALUNO = [("aluno_nome", "Nome do aluno", True), ("turma", "Turma", True), ("turno", "Turno", False), ("matricula", "Matrícula", False)]
CAMPOS_RESP = [("nome", "Nome"), ("parentesco", "Parentesco"), ("telefone", "Celular"), ("email", "E-mail")]

MODELO_CABECALHO = ["Aluno", "Turma", "Turno", "Matrícula"] + [
    f"{campo} {n}" for n in (1, 2, 3) for campo in ("Responsável", "Parentesco", "Celular", "E-mail")
]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
TURNOS = {"manha": "manha", "matutino": "manha", "m": "manha", "tarde": "tarde", "vespertino": "tarde", "t": "tarde",
          "integral": "integral", "i": "integral", "manha e tarde": "integral", "periodo integral": "integral"}
TURNO_ROTULO = {"manha": "manhã", "tarde": "tarde", "integral": "integral"}
PARTICULAS = {"da", "de", "do", "das", "dos", "e"}


class PlanilhaInvalida(ValueError):
    pass


# ---------------------------------------------------------------- leitura

def _norm(texto):
    """'Nome da Mãe ' → 'nome da mae'. Comparação de cabeçalho, turma e nome."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    t = re.sub(r"([a-z])(\d)", r"\1 \2", t)  # "resp1", "celular2"
    return " ".join(re.sub(r"[^a-z0-9]+", " ", t).split())


def _celula(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)  # celular digitado como número no Excel
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return " ".join(str(v).split())


def _ler_xlsx(conteudo):
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — arquivo corrompido/protegido: mensagem pra secretaria
        raise PlanilhaInvalida(f"não consegui abrir a planilha ({type(e).__name__}). Salve de novo como .xlsx ou CSV.") from e
    ws = wb["Alunos"] if "Alunos" in wb.sheetnames else wb.worksheets[0]
    tabela = []
    for linha in ws.iter_rows(values_only=True):
        tabela.append([_celula(v) for v in linha])
        if len(tabela) > MAX_LINHAS + 20:
            break
    wb.close()
    return tabela


def _ler_csv(conteudo):
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252", errors="replace")  # CSV salvo pelo Excel em português
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t|")
        delim = dialeto.delimiter
    except csv.Error:
        delim = ";" if texto.split("\n", 1)[0].count(";") else ","
    return [[_celula(c) for c in linha] for linha in csv.reader(io.StringIO(texto), delimiter=delim)]


def ler_arquivo(nome, conteudo):
    """bytes do upload → (cabecalho, linhas, linha_inicial). linha_inicial é o número da linha
    no Excel da primeira linha de dados, pra apontar "linha 37: sem turma"."""
    if not conteudo:
        raise PlanilhaInvalida("arquivo vazio")
    if len(conteudo) > MAX_BYTES:
        raise PlanilhaInvalida("arquivo maior que 5 MB")
    nome = (nome or "").lower()
    if nome.endswith(".xls") and conteudo[:2] != b"PK":
        raise PlanilhaInvalida("arquivo .xls antigo: no Excel use Arquivo › Salvar como › Pasta de Trabalho do Excel (.xlsx)")
    tabela = _ler_xlsx(conteudo) if conteudo[:2] == b"PK" else _ler_csv(conteudo)
    while tabela and not any(tabela[-1]):
        tabela.pop()
    for i, linha in enumerate(tabela):
        if sum(1 for c in linha if c) >= 2:  # cabeçalho = 1ª linha com 2+ células (pula título)
            largura = max(len(r) for r in tabela[i:])
            cab = [c or f"Coluna {k + 1}" for k, c in enumerate(linha + [""] * (largura - len(linha)))]
            dados = [r + [""] * (largura - len(r)) for r in tabela[i + 1:]]
            if not any(any(r) for r in dados):
                raise PlanilhaInvalida("a planilha só tem o cabeçalho — nenhum aluno preenchido")
            if len(dados) > MAX_LINHAS:
                raise PlanilhaInvalida(f"mais de {MAX_LINHAS} linhas: divida em arquivos menores")
            return cab, dados, i + 2
    raise PlanilhaInvalida("não encontrei o cabeçalho (linha com os nomes das colunas)")


# ---------------------------------------------------------------- mapeamento

IGNORAR = {"cpf", "rg", "documento", "nascimento", "endereco", "cep", "bairro", "cidade", "uf", "profissao", "sexo", "idade",
           "naturalidade", "nacionalidade", "observacao", "observacoes", "obs", "letivo", "situacao", "status", "data"}
PALAVRAS_RESP = {"responsavel", "responsaveis", "resp", "pais", "familiar", "contato"}
PALAVRAS_TEL = {"telefone", "celular", "fone", "tel", "whatsapp", "whats", "cel", "contato"}


def _classificar(h):
    """Cabeçalho normalizado → (grupo, campo) ou None. grupo: 'aluno', 'mae', 'pai' ou 'r1'..'r5'."""
    p = h.split()
    if not p or set(p) & IGNORAR:
        return None
    num = next((int(w) for w in p if w.isdigit() and 1 <= int(w) <= SLOTS), None)
    eh_email = "email" in p or "mail" in p
    eh_par = bool(set(p) & {"parentesco", "vinculo", "grau", "relacao"})
    eh_tel = bool(set(p) & PALAVRAS_TEL) and "nome" not in p
    if "mae" in p:
        grupo = "mae"
    elif "pai" in p:
        grupo = "pai"
    elif set(p) & PALAVRAS_RESP or eh_tel or eh_email or eh_par:
        grupo = f"r{num or 1}"
    else:
        if "turno" in p or "periodo" in p:
            return "aluno", "turno"
        if "turma" in p or "classe" in p or (set(p) & {"serie", "ano", "grupo", "sala"} and len(p) <= 3):
            return "aluno", "turma"
        if "matricula" in p or "codigo" in p or p in (["ra"], ["cod"], ["id"]):
            return "aluno", "matricula"
        if set(p) & {"aluno", "aluna", "estudante", "crianca"} or p == ["nome"] or ("nome" in p and "completo" in p):
            return "aluno", "aluno_nome"
        return None
    if eh_email:
        return grupo, "email"
    if eh_par:
        return grupo, "parentesco"
    if eh_tel:
        return grupo, "telefone"
    return grupo, "nome"


def adivinhar_mapeamento(cabecalho):
    """{'aluno_nome': idx, 'turma': idx, ..., 'resp': [{nome, parentesco, telefone, email, padrao}] * 5}"""
    mapa = {c: None for c, _, _ in CAMPOS_ALUNO}
    grupos, ordem = {}, []
    for i, h in enumerate(cabecalho):
        c = _classificar(_norm(h))
        if not c:
            continue
        grupo, campo = c
        if grupo == "aluno":
            if mapa[campo] is None:
                mapa[campo] = i
            continue
        if grupo not in grupos:
            grupos[grupo] = {}
            ordem.append(grupo)
        grupos[grupo].setdefault(campo, i)
    slots = [None] * SLOTS
    for grupo in ordem:  # "Responsável 2" vai pro slot 2; mãe e pai ocupam os livres
        if grupo.startswith("r") and slots[int(grupo[1:]) - 1] is None:
            slots[int(grupo[1:]) - 1] = grupo
    for grupo in ordem:
        if grupo not in slots:
            livre = next((k for k in range(SLOTS) if slots[k] is None), None)
            if livre is not None:
                slots[livre] = grupo
    mapa["resp"] = []
    for grupo in slots:
        campos = grupos.get(grupo, {}) if grupo else {}
        slot = {c: campos.get(c) for c, _ in CAMPOS_RESP}
        slot["padrao"] = {"mae": "mãe", "pai": "pai"}.get(grupo, "")
        mapa["resp"].append(slot)
    return mapa


def mapeamento_do_form(form, n_colunas):
    def idx(v):
        try:
            i = int(v)
        except (TypeError, ValueError):
            return None
        return i if 0 <= i < n_colunas else None

    mapa = {c: idx(form.get(f"m_{c}")) for c, _, _ in CAMPOS_ALUNO}
    mapa["resp"] = []
    for s in range(SLOTS):
        slot = {c: idx(form.get(f"m_r{s}_{c}")) for c, _ in CAMPOS_RESP}
        slot["padrao"] = (form.get(f"m_r{s}_padrao") or "").strip().lower()[:40]
        mapa["resp"].append(slot)
    return mapa


def faltando(mapa):
    """Campos obrigatórios sem coluna escolhida."""
    return [rotulo for c, rotulo, obrig in CAMPOS_ALUNO if obrig and mapa.get(c) is None]


# ---------------------------------------------------------------- linhas → registros

def _nome(v):
    """'MARIA DA SILVA' (export de sistema) → 'Maria da Silva'."""
    v = " ".join((v or "").split())
    if v.isupper():
        v = " ".join(w.lower() if w.lower() in PARTICULAS and k else w.capitalize() for k, w in enumerate(v.lower().split()))
    return v


def _valor(linha, idx):
    return (linha[idx] if idx is not None and idx < len(linha) else "").strip()


def turno_de(texto):
    return TURNOS.get(_norm(texto))


def _mesmo_resp(a, b):
    if a["telefone"] and b["telefone"]:
        return mesmo_telefone(a["telefone"], b["telefone"])
    if a["email"] and b["email"] and a["email"] == b["email"]:
        return True
    return _norm(a["nome"]) == _norm(b["nome"])


def processar(linhas, mapa, linha_inicial=2):
    """Aplica o mapeamento. Linhas do mesmo aluno (formato "uma linha por responsável") se
    juntam. Devolve registros {linhas, aluno, responsaveis, erros, avisos}."""
    registros, por_chave = [], {}
    for k, linha in enumerate(linhas):
        n = linha_inicial + k
        if not any((c or "").strip() for c in linha):
            continue
        nome = _nome(_valor(linha, mapa.get("aluno_nome")))
        turma = " ".join(_valor(linha, mapa.get("turma")).split())
        matricula = _valor(linha, mapa.get("matricula"))
        turno_txt = _valor(linha, mapa.get("turno"))
        erros, avisos = [], []
        if not nome:
            erros.append("sem nome do aluno")
        if not turma:
            erros.append("sem turma")
        turno = turno_de(turno_txt) if turno_txt else None
        if turno_txt and not turno:
            avisos.append(f"turno “{turno_txt}” não reconhecido (use manhã, tarde ou integral)")
        resps = []
        for s, slot in enumerate(mapa.get("resp") or []):
            rnome = _nome(_valor(linha, slot.get("nome")))
            tel = _valor(linha, slot.get("telefone"))
            email = _valor(linha, slot.get("email")).lower()
            if not (rnome or tel or email):
                continue
            if not rnome:
                avisos.append(f"responsável {s + 1} sem nome — ignorado")
                continue
            if tel and len(normalizar_telefone(tel)) not in (10, 11):
                avisos.append(f"celular de {rnome} inválido ({tel}) — importado sem celular")
                tel = ""
            if email and not EMAIL_RE.match(email):
                avisos.append(f"e-mail de {rnome} inválido ({email}) — importado sem e-mail")
                email = ""
            parentesco = (_valor(linha, slot.get("parentesco")) or slot.get("padrao") or "").strip().lower()[:40]
            resps.append({"nome": rnome, "telefone": normalizar_telefone(tel), "email": email, "parentesco": parentesco})
        if erros:
            registros.append({"linhas": [n], "aluno": {"nome": nome, "turma": turma}, "responsaveis": [], "erros": erros, "avisos": avisos})
            continue
        chave = ("m", matricula.lower()) if matricula else ("n", _norm(nome), _norm(turma))
        reg = por_chave.get(chave)
        if reg is None:
            reg = {"linhas": [n], "aluno": {"nome": nome, "turma": turma, "turno": turno, "matricula": matricula},
                   "responsaveis": [], "erros": [], "avisos": []}
            por_chave[chave] = reg
            registros.append(reg)
        else:
            reg["linhas"].append(n)
            reg["aluno"]["turno"] = reg["aluno"]["turno"] or turno
        reg["avisos"] += avisos
        for r in resps:
            igual = next((x for x in reg["responsaveis"] if _mesmo_resp(r, x)), None)
            if igual:
                for campo in ("telefone", "email", "parentesco"):
                    igual[campo] = igual[campo] or r[campo]
            else:
                reg["responsaveis"].append(r)
    for reg in registros:
        if len(reg["responsaveis"]) > SLOTS:
            reg["avisos"].append(f"mais de {SLOTS} responsáveis — importados só os {SLOTS} primeiros")
            del reg["responsaveis"][SLOTS:]
    return registros


# ---------------------------------------------------------------- plano (prévia) e execução

def _chave_tel(t):
    t = normalizar_telefone(t)
    return (t[:2], t[-8:]) if len(t) >= 10 else None


def planejar(unidade_id, registros):
    """O que vai acontecer, sem gravar nada. A mesma função roda de novo dentro da transação
    na hora de importar, então a prévia é exatamente o que será gravado."""
    turmas = {_norm(t["nome"]): t for t in db.query("SELECT id, nome, turno FROM turmas WHERE unidade_id = %s", (unidade_id,))}
    alunos = db.query(
        "SELECT a.id, a.pessoa_id, a.turma_id, a.matricula, p.nome FROM alunos a JOIN pessoas p ON p.id = a.pessoa_id "
        "WHERE p.unidade_id = %s AND p.status <> 'desligado'", (unidade_id,),
    )
    por_matricula = {a["matricula"].strip().lower(): a for a in alunos if a["matricula"]}
    por_nome_turma = {(_norm(a["nome"]), a["turma_id"]): a for a in alunos}
    resps = db.query(
        "SELECT r.id, r.telefone, r.email, p.nome FROM responsaveis r JOIN pessoas p ON p.id = r.pessoa_id "
        "WHERE p.unidade_id = %s AND p.status <> 'desligado'", (unidade_id,),
    )
    por_tel = {}
    for r in resps:
        if _chave_tel(r["telefone"]):
            por_tel.setdefault(_chave_tel(r["telefone"]), r)
    por_email = {r["email"].lower(): r for r in resps if r["email"]}
    resp_por_id = {r["id"]: r for r in resps}
    vinculos = {(v["aluno_id"], v["responsavel_id"]) for v in db.query(
        "SELECT ar.aluno_id, ar.responsavel_id FROM aluno_responsavel ar JOIN alunos a ON a.id = ar.aluno_id "
        "JOIN pessoas p ON p.id = a.pessoa_id WHERE p.unidade_id = %s", (unidade_id,))}
    n_vinculos = Counter(a for a, _ in vinculos)

    plano = {"turmas_novas": {}, "responsaveis_novos": [], "itens": [], "erros": [], "avisos": []}
    existentes_usados = set()
    for reg in registros:
        if reg["erros"]:
            plano["erros"].append({"linhas": reg["linhas"], "aluno": reg["aluno"]["nome"], "erros": reg["erros"]})
            continue
        for a in reg["avisos"]:
            plano["avisos"].append({"linhas": reg["linhas"], "aluno": reg["aluno"]["nome"], "aviso": a})
        al = reg["aluno"]
        chave_turma = _norm(al["turma"])
        if chave_turma in turmas:
            turma_ref, turma_id = ("id", turmas[chave_turma]["id"]), turmas[chave_turma]["id"]
        else:
            nova = plano["turmas_novas"].setdefault(chave_turma, {"nome": al["turma"], "turno": None})
            nova["turno"] = nova["turno"] or al["turno"]
            turma_ref, turma_id = ("nova", chave_turma), None
        existente = por_matricula.get(al["matricula"].lower()) if al["matricula"] else None
        if not existente and turma_id:
            existente = por_nome_turma.get((_norm(al["nome"]), turma_id))
        item = {"linhas": reg["linhas"], "aluno_nome": al["nome"], "turma_nome": al["turma"], "matricula": al["matricula"],
                "turma": turma_ref, "aluno": existente, "muda_turma": bool(existente and existente["turma_id"] != turma_id),
                "responsaveis": []}
        ja_tem = n_vinculos[existente["id"]] if existente else 0
        for r in reg["responsaveis"]:
            achado = por_tel.get(_chave_tel(r["telefone"])) if r["telefone"] else None
            if not achado and r["email"]:
                achado = por_email.get(r["email"])
            if not achado and existente:  # mesmo nome já vinculado a este aluno (planilha sem celular)
                achado = next((resp_por_id[rid] for (aid, rid) in vinculos if aid == existente["id"] and _norm(resp_por_id.get(rid, {}).get("nome")) == _norm(r["nome"])), None)
            if achado:
                ref = ("id", achado["id"])
                existentes_usados.add(achado["id"])
                completar = {c: r[c] for c in ("telefone", "email") if r[c] and not achado[c]}
                novo_vinculo = not existente or (existente["id"], achado["id"]) not in vinculos
            else:
                k = next((i for i, x in enumerate(plano["responsaveis_novos"]) if _mesmo_resp(r, x)), None)
                if k is None:
                    plano["responsaveis_novos"].append(dict(r))
                    k = len(plano["responsaveis_novos"]) - 1
                ref, completar, novo_vinculo = ("novo", k), {}, True
            if any(x["ref"] == ref for x in item["responsaveis"]):
                continue
            if novo_vinculo and ja_tem >= SLOTS:
                plano["avisos"].append({"linhas": reg["linhas"], "aluno": al["nome"], "aviso": f"{r['nome']} não vinculado: aluno já tem {SLOTS} responsáveis"})
                continue
            ja_tem += 1 if novo_vinculo else 0
            item["responsaveis"].append({"ref": ref, "nome": r["nome"], "telefone": r["telefone"], "email": r["email"],
                                         "parentesco": r["parentesco"], "completar": completar, "vinculo_novo": novo_vinculo})
        plano["itens"].append(item)
    for t in plano["turmas_novas"].values():
        t["turno"] = t["turno"] or "manha"
    itens = plano["itens"]
    plano["resumo"] = {
        "alunos_novos": sum(1 for i in itens if not i["aluno"]),
        "alunos_existentes": sum(1 for i in itens if i["aluno"]),
        "mudam_turma": sum(1 for i in itens if i["muda_turma"]),
        "turmas_novas": [f"{t['nome']} ({TURNO_ROTULO[t['turno']]})" for t in plano["turmas_novas"].values()],
        "responsaveis_novos": len(plano["responsaveis_novos"]),
        "responsaveis_existentes": len(existentes_usados),
        "vinculos_novos": sum(1 for i in itens for r in i["responsaveis"] if r["vinculo_novo"]),
        "linhas_com_erro": len(plano["erros"]),
        "avisos": len(plano["avisos"]),
        "sem_responsavel": sum(1 for i in itens if not i["responsaveis"] and not i["aluno"]),
    }
    return plano


def executar(unidade_id, registros, consentimento=None):
    """Grava o plano numa transação só. `consentimento` = {assinado_por, observacao} registra
    o termo LGPD de cada pessoa NOVA (a escola declara que já tem o termo da matrícula)."""
    with db.lote():
        plano = planejar(unidade_id, registros)
        ids_turma = {}
        for chave, t in plano["turmas_novas"].items():
            ids_turma[chave] = db.execute("INSERT INTO turmas (unidade_id, nome, turno) VALUES (%s,%s,%s) RETURNING id",
                                          (unidade_id, t["nome"], t["turno"]), returning=True)["id"]
            auditar("criar", "turma", ids_turma[chave], depois=t, unidade_id=unidade_id)
        novas_pessoas, ids_resp = [], []
        for r in plano["responsaveis_novos"]:
            pessoa_id, resp_id = service.criar_responsavel(unidade_id, r["nome"], r["telefone"] or None)
            if r["email"]:
                db.execute("UPDATE responsaveis SET email = %s WHERE id = %s", (r["email"], resp_id))
            ids_resp.append(resp_id)
            novas_pessoas.append(pessoa_id)
        for item in plano["itens"]:
            turma_id = item["turma"][1] if item["turma"][0] == "id" else ids_turma[item["turma"][1]]
            if item["aluno"]:
                aluno_id = item["aluno"]["id"]
                if item["muda_turma"]:
                    db.execute("UPDATE alunos SET turma_id = %s WHERE id = %s", (turma_id, aluno_id))
                    auditar("editar", "aluno", aluno_id, antes={"turma_id": item["aluno"]["turma_id"]}, depois={"turma_id": turma_id}, unidade_id=unidade_id)
                    dispositivos.enfileirar_pessoa(item["aluno"]["pessoa_id"], unidade_id, com_foto=False)  # regra de saída depende da turma
            else:
                pessoa_id, aluno_id = service.criar_aluno(unidade_id, item["aluno_nome"], turma_id, item["matricula"] or None)
                novas_pessoas.append(pessoa_id)
            for r in item["responsaveis"]:
                resp_id = r["ref"][1] if r["ref"][0] == "id" else ids_resp[r["ref"][1]]
                if r["completar"]:
                    db.execute("UPDATE responsaveis SET telefone = coalesce(telefone, %s), email = coalesce(email, %s) WHERE id = %s",
                               (r["completar"].get("telefone"), r["completar"].get("email"), resp_id))
                if r["vinculo_novo"]:
                    service.vincular(aluno_id, resp_id, r["parentesco"] or None)
        if consentimento:
            for pessoa_id in novas_pessoas:
                service.registrar_consentimento(pessoa_id, consentimento["assinado_por"], date.today(), observacao=consentimento.get("observacao"),
                                                canal="importacao")
    return plano["resumo"]


# ---------------------------------------------------------------- modelo para a escola

INSTRUCOES = [
    "Como preencher (uma linha por aluno)",
    "",
    "1. Obrigatórios: Aluno e Turma. Use o nome da turma como a escola chama (ex.: Pré 1, 3º ano A). Turma que não existir no sistema é criada.",
    "2. Turno: manhã, tarde ou integral. Só é usado quando a turma ainda não existe.",
    "3. Irmãos: uma linha para cada um, repetindo os responsáveis. O sistema junta a mesma mãe/pai pelo celular.",
    "4. Celular com DDD, ex.: (21) 99999-0000. É por ele que o responsável ativa os avisos no Telegram.",
    "5. E-mail: para a escola mandar o convite dos avisos por e-mail.",
    "6. Mais de 3 responsáveis? Acrescente as colunas Responsável 4, Parentesco 4, Celular 4, E-mail 4 (até 5).",
    "7. Não precisa usar este modelo: a planilha do sistema da escola também serve. Na hora de importar você confere as colunas.",
    "8. A aba Exemplo é só para consulta e não é importada.",
]
EXEMPLO = [
    ["Laura Mendes", "Pré 1", "manhã", "2026-031", "Carla Mendes", "mãe", "(21) 98888-0001", "carla@exemplo.com", "Rafael Mendes", "pai", "(21) 97777-0002", "", "", "", "", ""],
    ["Pedro Mendes", "2º ano", "tarde", "2026-032", "Carla Mendes", "mãe", "(21) 98888-0001", "carla@exemplo.com", "Sônia Ribeiro", "avó", "(21) 96666-0003", "", "", "", "", ""],
]


def modelo_xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alunos"
    for aba, linhas in ((ws, []), (wb.create_sheet("Exemplo"), EXEMPLO)):
        aba.append(MODELO_CABECALHO)
        for linha in linhas:
            aba.append(linha)
        for k, cel in enumerate(aba[1]):
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="2563EB" if k < 2 else ("1E293B" if k < 4 else "475569"))
            aba.column_dimensions[cel.column_letter].width = 26 if k in (0, 4, 8, 12) else 16
        aba.freeze_panes = "A2"
    turno = DataValidation(type="list", formula1='"manhã,tarde,integral"', allow_blank=True)
    ws.add_data_validation(turno)
    turno.add("C2:C2000")
    for col in ("D", "G", "K", "O"):  # matrícula e celulares como texto: Excel não come o zero nem vira 2,19E+10
        for linha in range(2, 1001):
            ws[f"{col}{linha}"].number_format = "@"
    ins = wb.create_sheet("Instruções")
    for linha in INSTRUCOES:
        ins.append([linha])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
